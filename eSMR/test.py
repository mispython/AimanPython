#!/usr/bin/env python3
"""
Program : EIBDDCMG.py
Purpose : Store Daily Position for EBANK/WALKER/DEPOSIT/EQUATION
          (Bank's Total Deposits - RM'000, gold GIA balance, NID/PMMD
          proceeds, and FX rates), appended into a running monthly
          DCMG dataset and printed as a month-to-date listing (Report ID: DMMISR1G).

eSMR 4294 : Exclude post-dated PMMD/PQMMD placements (STRTDATE > REPTDATE)
            from TOTCPMMD / TOTIPMMD / TOTNPMMD, and show them in 3 new
            columns (PDCPMMD / PDIPMMD / PDNPMMD) placed before CLOSBAL.
            DFI & FI columns (TOTCDFI / TOTIDFI / TOTNDFI) are UNCHANGED.

Report date : BATCH_DTTM of SOURCE_SYSTEM_CD = 'DP' in ctl_dwh_batch_dttm.sas7bdat.

Monthly rows : each run stores the report-date row, then recomputes the 6 PMMD
               columns (3 existing + 3 post-dated) of the earlier dates of the
               month from each day's own UTFX / IUTFX file, so every row of the
               month is on the same basis. If a daily file is missing, that
               row is left as saved and a warning is printed.

Parquet     : each .sas7bdat input is converted to a TEMPORARY parquet (per-run
               temp folder), read from it, then deleted. Nothing is kept.

Output      : one-table Excel (.xlsx) and one-table text (.TXT)
SFTP        : the .TXT is sent to DRR (JCL STEP01 + RUNSFTP).
"""

import gc
import csv
import math
from pathlib import Path
from datetime import date, datetime, timedelta
from typing import Optional

import atexit
import shutil
import sys
import tempfile
import paramiko

import duckdb
import pandas as pd
import polars as pl
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ============================================================================
# PATH CONFIGURATION  (folders)
# ============================================================================
# Production Path
INPUT_DIR   = Path("/dwh")
RNID_DIR    = INPUT_DIR / "rnid"                            # RNID (.sas7bdat)
IRNID_DIR   = INPUT_DIR / "irnid"                           # IRNID (.sas7bdat)
DEQ_DIR     = INPUT_DIR / "eq_d"                            # UTFX (.sas7bdat)
IDEQ_DIR    = INPUT_DIR / "ieq_d"                           # IUTFX (.sas7bdat)

HOST_DIR    = Path("/host_pq/mis")
TEXT_DIR    = HOST_DIR / "input/textfile"                   # DPGIA, EQ, IEQ (.txt)
EGOLD_DIR   = Path("/extdwh/host")                                     # EFORATE + legacy DCMG master (.sas7bdat)
STORE_DIR   = HOST_DIR / "parquet/EIBDDCMG"                 # Monthly DCMG masters (parquet)
OUTPUT_DIR  = HOST_DIR / "output/deposit/report"            # Report output

for _d in (STORE_DIR, OUTPUT_DIR):
    _d.mkdir(parents=True, exist_ok=True)

# # Testing Path
# BASE_DIR   = Path("/sas/python/virt_edw/Data_Warehouse/MIS/XMIS")

# INPUT_DIR  = BASE_DIR / "input"  / "prod"  / "EIBDDCMG"     # RNID, IRNID, UTFX, IUTFX (.sas7bdat)
# HOST_DIR   = Path("/stgsrcsys/host/uat/AII/EIBDDCMG")       # DPGIA, EQ, IEQ (.txt)
# EGOLD_DIR  = HOST_DIR                                       # EFORATE + legacy DCMG master (.sas7bdat)
# CACHE_DIR  = BASE_DIR / "input"  / "cache" / "EIBDDCMG"     # Parquet cache for .sas7bdat sources
# STORE_DIR  = BASE_DIR / "store"  / "EIBDDCMG"               # Monthly DCMG masters (parquet)
# OUTPUT_DIR = BASE_DIR / "output" / "EIBDDCMG"               # Report output

# for _d in (CACHE_DIR, STORE_DIR, OUTPUT_DIR):
#     _d.mkdir(parents=True, exist_ok=True)

CHUNK_ROWS = 500_000

# Temporary parquet area: created per run, removed on exit (nothing is kept)
TEMP_DIR  = None   # None = system temp; set a folder if the default temp area is too small
CACHE_DIR = Path(tempfile.mkdtemp(prefix="EIBDDCMG_", dir=TEMP_DIR))
atexit.register(shutil.rmtree, CACHE_DIR, ignore_errors=True)

# Control table that holds the report date  -- TODO: set the real folder
INPUT_CTL_BATCH_FILE = Path("/sasdata/dwh/control/ctl_dwh_batch_dttm.sas7bdat")

# CUSTFISS codes treated as DFI & FI (unchanged from original)
DFI_CODES = ('02', '03', '07', '10', '12', '81', '82', '83', '84', '01')

# ============================================================================
# HELPERS: CACHE (.sas7bdat -> Parquet)   [same pattern as EIBMRB01.py]
# ============================================================================
# Unused-1 : freshness check no longer needed (parquet is temporary, never reused across runs)
# def _cache_is_fresh(sas_path: Path, cache_path: Path) -> bool:
#     return cache_path.exists() and cache_path.stat().st_mtime >= sas_path.stat().st_mtime


def _sas_type_to_arrow(dtype) -> pa.DataType:
    if dtype == "object" or pd.api.types.is_string_dtype(dtype):
        return pa.string()
    if pd.api.types.is_integer_dtype(dtype):
        return pa.int64()
    if pd.api.types.is_float_dtype(dtype):
        return pa.float64()
    return pa.from_numpy_dtype(dtype)


def _sas_to_parquet(sas_path: Path, cache_path: Path, tag: str) -> None:
    print(f"  [{tag}] Converting {sas_path.name} -> {cache_path.name} ...")
    writer = None
    schema = None
    total = 0
    for chunk in pd.read_sas(sas_path, encoding="latin1", chunksize=CHUNK_ROWS):
        if schema is None:
            schema = pa.schema([pa.field(c, _sas_type_to_arrow(t)) for c, t in chunk.dtypes.items()])
            writer = pq.ParquetWriter(cache_path, schema, compression="snappy")
        table = pa.Table.from_pandas(chunk, schema=schema, preserve_index=False)
        writer.write_table(table)
        total += len(chunk)
        del chunk, table
        gc.collect()
    if writer:
        writer.close()
    print(f"  [{tag}] Done - {total:,} rows cached.")


# Unused-2 : old persistent-cache version of _load_cached (replaced by the temporary-parquet version below)
# def _load_cached(sas_path: Path, tag: str) -> Path:
#     # Cache name = parent folder + file stem, so same-named files in different
#     # folders never share a cache.
#     if not sas_path.exists():
#         raise FileNotFoundError(f"[{tag}] input not found: {sas_path}")
#     cache_path = CACHE_DIR / f"{sas_path.parent.name}_{sas_path.stem}.parquet"
#     if _cache_is_fresh(sas_path, cache_path):
#         print(f"  [{tag}] Cache fresh - skipping conversion.")
#     else:
#         _sas_to_parquet(sas_path, cache_path, tag)
#     return cache_path


def _load_cached(sas_path: Path, tag: str) -> Path:
    # Converts .sas7bdat -> temporary Parquet (parent folder + stem = unique name).
    # The file lives only in this run's temp folder.
    if not sas_path.exists():
        raise FileNotFoundError(f"[{tag}] input not found: {sas_path}")
    cache_path = CACHE_DIR / f"{sas_path.parent.name}_{sas_path.stem}.parquet"
    if cache_path.exists():                      # already converted earlier in this run
        return cache_path
    try:
        _sas_to_parquet(sas_path, cache_path, tag)
    except BaseException:
        cache_path.unlink(missing_ok=True)       # never leave a half-written parquet
        raise
    return cache_path


def _discard(*paths) -> None:
    """Delete temporary parquet file(s) as soon as they have been read."""
    for p in paths:
        Path(p).unlink(missing_ok=True)

# ============================================================================
# STEP 1: REPORT DATE  (ctl_dwh_batch_dttm, SOURCE_SYSTEM_CD = 'DP')
# ============================================================================
print("Step 1: Deriving report date...")

def _read_dp_batch_date() -> date:
    cache = _load_cached(INPUT_CTL_BATCH_FILE, "CTL_BATCH")
    con = duckdb.connect(database=":memory:")
    try:
        r = con.execute(
            f"SELECT BATCH_DTTM FROM read_parquet('{cache.as_posix()}') "
            f"WHERE TRIM(SOURCE_SYSTEM_CD) = 'DP'"
        ).fetchone()
    finally:
        con.close()
    if r is None or r[0] is None:
        raise ValueError("SOURCE_SYSTEM_CD = 'DP' not found in ctl_dwh_batch_dttm")
    v = r[0]
    if isinstance(v, datetime):          # check before date (datetime is a date subclass)
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):               # e.g. '23AUG26:00:00:00'
        return datetime.strptime(v.strip()[:7].upper(), "%d%b%y").date()
    return (datetime(1960, 1, 1) + timedelta(seconds=float(v))).date()   # raw SAS datetime

reptdate       = _read_dp_batch_date()
# reptdate       = date(2026, 10, 19)
REPTYEAR       = reptdate.strftime("%y")           # 2-digit year
REPTMON        = reptdate.strftime("%m")
REPTDAY        = reptdate.strftime("%d")
RPDATE         = reptdate.strftime("%d/%m/%y")     # DDMMYY8.
RDATE_DDMMYY10 = reptdate.strftime("%d/%m/%Y")     # DDMMYY10.

print(f"  Report date : {RDATE_DDMMYY10}  (RPDATE={RPDATE})")
print(f"  REPTMON={REPTMON}  REPTDAY={REPTDAY}  REPTYEAR={REPTYEAR}")

# ============================================================================
# PHYSICAL INPUTS  (one line per file; change each path independently)
# ============================================================================
# --- text files, GDG (0) = latest generation, fixed name ---
INPUT_EQ_FILE    = TEXT_DIR / "PBB_EQNID.txt"       # //EQ    DD SAP.PBB.EQNID.TXT(0)
INPUT_IEQ_FILE   = TEXT_DIR / "PIBB_EQNID.txt"      # //IEQ   DD SAP.PIBB.EQNID.TXT(0)
INPUT_DPGIA_FILE = TEXT_DIR / "DPGIA.txt"           # //DPGIA DD RBP2.B033.PBGIA.DLYPOS.RPT(0)

# # --- text files, GDG (0) = latest generation, fixed name ---
# INPUT_EQ_FILE    = HOST_DIR / "PBB_EQNID.txt"       # //EQ    DD SAP.PBB.EQNID.TXT(0)
# INPUT_IEQ_FILE   = HOST_DIR / "PIBB_EQNID.txt"      # //IEQ   DD SAP.PIBB.EQNID.TXT(0)
# INPUT_DPGIA_FILE = HOST_DIR / "DPGIA.txt"           # //DPGIA DD RBP2.B033.PBGIA.DLYPOS.RPT(0)

# --- sas7bdat, file name carries the report date ---
# Unused-3 : UTFX_DIR / IUTFX_DIR are not defined anywhere (NameError); replaced by DEQ_DIR / IDEQ_DIR
# INPUT_UTFX_DIR   = UTFX_DIR                        # //DEQ   DD SAP.PBB.EQUT(0)
# INPUT_IUTFX_DIR  = IUTFX_DIR                       # //IDEQ  DD SAP.PIBB.EQUT(0)
INPUT_UTFX_DIR   = DEQ_DIR                         # //DEQ   DD SAP.PBB.EQUT(0)
INPUT_IUTFX_DIR  = IDEQ_DIR                        # //IDEQ  DD SAP.PIBB.EQUT(0)
INPUT_RNID_DIR   = RNID_DIR                        # //RNID  DD SAP.PBB.RNID.SASDATA
INPUT_IRNID_DIR  = IRNID_DIR                       # //IRNID DD SAP.PIBB.RNID.SASDATA

def _utfx_file(d: date)  -> Path: return INPUT_UTFX_DIR  / f"utfx{d:%y%m%d}.sas7bdat"
def _iutfx_file(d: date) -> Path: return INPUT_IUTFX_DIR / f"iutfx{d:%y%m%d}.sas7bdat"

INPUT_UTFX_FILE  = _utfx_file(reptdate)             # SAS: DEQ.UTFX&REPTYEAR&REPTMON&REPTDAY
INPUT_IUTFX_FILE = _iutfx_file(reptdate)            # SAS: IDEQ.IUTFX&REPTYEAR&REPTMON&REPTDAY
INPUT_RNID_FILE  = INPUT_RNID_DIR  / f"rnid{REPTYEAR}{REPTMON}{REPTDAY}.sas7bdat"    # SAS: RNID.RNID&REPTDAY
INPUT_IRNID_FILE = INPUT_IRNID_DIR / f"irnid{REPTYEAR}{REPTMON}{REPTDAY}.sas7bdat"   # SAS: IRNID.RNID&REPTDAY

# --- sas7bdat, file name carries the report month ---
INPUT_EFORATE_FILE     = EGOLD_DIR / f"eforate{REPTMON}.sas7bdat"   # MIS.EFORATE&REPTMON
INPUT_DCMG_LEGACY_FILE = EGOLD_DIR / f"dcmg{REPTMON}.sas7bdat"      # MIS.DCMG&REPTMON (bootstrap only)

# ============================================================================
# PHYSICAL OUTPUTS  (original: DMMISR1G@%OMM.%OYY..TXT -> DMMISR1G@{MM}{YY}.TXT)
# ============================================================================
OUTPUT_TXT_FILE  = OUTPUT_DIR / f"DMMISR1G@{REPTMON}{REPTYEAR}.TXT"
OUTPUT_XLSX_FILE = OUTPUT_DIR / f"DMMISR1G@{REPTMON}{REPTYEAR}.xlsx"
STORE_DCMG_FILE  = STORE_DIR  / f"dcmg{REPTMON}.parquet"            # monthly master (read and written)

# ============================================================================
# HELPERS: NUMERIC PARSING / SAS FUNCTIONS
# ============================================================================
def _parse_comma_numeric(text) -> Optional[float]:
    """Parse a COMMA-formatted numeric text field."""
    if text is None:
        return None
    s = str(text).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _parse_w_d_numeric(text: str, decimals: int) -> Optional[float]:
    """Plain SAS w.d informat: explicit decimal point wins, else implied decimals."""
    s = text.strip()
    if not s:
        return None
    try:
        if "." in s:
            return float(s)
        return int(s) / (10 ** decimals)
    except ValueError:
        return None


def _parse_ddmmyyyy(text: str) -> Optional[date]:
    s = text.strip()
    if len(s) < 10:
        return None
    day, month, year = s[0:2], s[3:5], s[6:10]
    try:
        return date(int(year), int(month), int(day))
    except (ValueError, TypeError):
        return None


def _sas_round(value, round_to: int = 1000):
    """SAS ROUND(): half away from zero."""
    if value is None:
        return None
    scaled = value / round_to
    rounded = math.floor(scaled + 0.5) if scaled >= 0 else math.ceil(scaled - 0.5)
    return rounded * round_to


def _thousands(value):
    """INT(ROUND(value,1000)/1000) -- RM'000, truncated."""
    r = _sas_round(value, 1000)
    if r is None:
        return None
    return int(r / 1000)


def _sas_sum(*values):
    """SAS SUM(): ignores missing; None only if all missing."""
    present = [v for v in values if v is not None]
    if not present:
        return None
    return sum(present)


def _fmt_comma(value, width: int, decimals: int = 0) -> str:
    """COMMAw.d display; OPTIONS MISSING=0 -> missing shows as 0."""
    if value is None:
        return "0".rjust(width)
    if decimals > 0:
        s = f"{value:,.{decimals}f}"
    else:
        s = f"{int(round(value)):,}"
    return s.rjust(width)


def _fmt_num(value, width: int, decimals: int = 2) -> str:
    if value is None:
        return "0".rjust(width)
    return f"{value:,.{decimals}f}".rjust(width)

# ============================================================================
# STEP 2: UTMS  (EQ)
# ============================================================================
print("\nStep 2: Reading UTMS (EQ)...")

def _read_pipe_delimited(path: Path) -> list[list[str]]:
    with open(path, "r", encoding="latin1", newline="") as fh:
        return list(csv.reader(fh, delimiter="|"))

_eq_rows = _read_pipe_delimited(INPUT_EQ_FILE)

utms_rows = []
for fields in _eq_rows[1:]:                      # FIRSTOBS=2
    proceed = _parse_comma_numeric(fields[2]) if len(fields) > 2 else None
    utms_rows.append({"REPTDATE": reptdate, "PROCEED": proceed})

utms = (
    pl.DataFrame(utms_rows, schema={"REPTDATE": pl.Date, "PROCEED": pl.Float64})
    .group_by("REPTDATE")
    .agg(pl.col("PROCEED").sum())
)
print(f"  UTMS rows: {len(utms):,}")

# ============================================================================
# STEP 3: INONFI + IFI -> IUTMS  (IEQ)
# ============================================================================
print("\nStep 3: Reading IUTMS (IEQ)...")

_ieq_rows = _read_pipe_delimited(INPUT_IEQ_FILE)

# INONFI: FIRSTOBS=2 OBS=2 -> only the second physical line
inonfi_rows = []
if len(_ieq_rows) > 1:
    fields = _ieq_rows[1]
    iproceed = _parse_comma_numeric(fields[2]) if len(fields) > 2 else None
    inonfi_rows.append({"REPTDATE": reptdate, "IPROCEED": iproceed, "IPROCEFI": None})

# IFI: FIRSTOBS=3 -> every line from the third onward
ifi_rows = []
for fields in _ieq_rows[2:]:
    iprocefi = _parse_comma_numeric(fields[2]) if len(fields) > 2 else None
    ifi_rows.append({"REPTDATE": reptdate, "IPROCEED": None, "IPROCEFI": iprocefi})

iutms_schema = {"REPTDATE": pl.Date, "IPROCEED": pl.Float64, "IPROCEFI": pl.Float64}
iutms = pl.concat(
    [
        pl.DataFrame(inonfi_rows, schema=iutms_schema),
        pl.DataFrame(ifi_rows, schema=iutms_schema),
    ],
    how="vertical",
).group_by("REPTDATE").agg(
    pl.col("IPROCEED").sum(),
    pl.col("IPROCEFI").sum(),
)
print(f"  IUTMS rows: {len(iutms):,}")

# ============================================================================
# STEP 4: EQUTMS = MERGE UTMS IUTMS BY REPTDATE
# ============================================================================
equtms = utms.join(iutms, on="REPTDATE", how="full", coalesce=True)

# ============================================================================
# STEP 5: PBGIA  (DPGIA)
# Only the 2nd balance line after each date line (N=3) is kept.
# ============================================================================
print("\nStep 5: Parsing PBGIA (DPGIA)...")

def _parse_pbgia(path: Path) -> pl.DataFrame:
    with open(path, "r", encoding="latin1") as fh:
        all_lines = fh.readlines()

    n = 0
    block_reptdate = None
    closbal = None
    records = []

    for raw in all_lines[7:]:                      # FIRSTOBS=8
        line = raw.rstrip("\r\n")
        if len(line) < 110:
            continue

        check1 = _parse_ddmmyyyy(line[1:11])       # @002 CHECK1 DDMMYY10.
        check2 = _parse_w_d_numeric(line[94:110], 2)   # @095 CHECK2 16.2

        if check1 is not None:
            n = 1
        elif check2 is not None:
            n += 1

        if check2 is None:                         # IF CHECK2 EQ . THEN DELETE
            continue

        if n == 1:
            block_reptdate = _parse_ddmmyyyy(line[1:11])
        else:
            closbal = check2

        if n == 3:
            records.append({"REPTDATE": block_reptdate, "CLOSBAL": closbal})

    return pl.DataFrame(records, schema={"REPTDATE": pl.Date, "CLOSBAL": pl.Float64})

pbgia = _parse_pbgia(INPUT_DPGIA_FILE)
print(f"  PBGIA rows (all dates in report): {len(pbgia):,}")
pbgia_today = pbgia.filter(pl.col("REPTDATE") == reptdate)

# ============================================================================
# STEP 6: UTFX / IUTFX  -- eSMR 4294 CHANGE
#   Existing PMMD columns : STRTDATE <= REPTDATE  (or STRTDATE missing)
#   New post-dated columns: STRTDATE >  REPTDATE
#   Both use the same non-DFI CUSTFISS rule.
#   DFI & FI columns unchanged (all deals, regardless of STRTDATE).
# ============================================================================
print("\nStep 6: Building UTFX / IUTFX (DEQ / IDEQ)...")

def _startdate_sql(cache_path: Path) -> str:
    """
    SQL expression returning STRTDATE as a DATE.
    pd.read_sas() gives a datetime only when the SAS column carries a date
    format; otherwise a raw SAS number (days since 1960-01-01).
    """
    con = duckdb.connect(database=":memory:")
    try:
        rows = con.execute(
            f"DESCRIBE SELECT STRTDATE FROM read_parquet('{cache_path.as_posix()}')"
        ).fetchall()
    finally:
        con.close()
    col_type = str(rows[0][1]).upper()
    if col_type.startswith(("DATE", "TIMESTAMP")):
        return "CAST(STRTDATE AS DATE)"
    return "(DATE '1960-01-01' + TRY_CAST(STRTDATE AS INTEGER))"


def _sum_deals(cache_path: Path, deal_types, rd: date,
               k_main: str, k_dfi: str, k_pd: str) -> dict:
    dt  = ",".join(f"'{x}'" for x in deal_types)
    dfi = ",".join(f"'{x}'" for x in DFI_CODES)
    sd  = _startdate_sql(cache_path)
    con = duckdb.connect(database=":memory:")
    try:
        r = con.execute(f"""
            WITH base AS (
                SELECT
                    CAST(AMTPAY AS DOUBLE) AS AMTPAY,
                    CUSTFISS,
                    COALESCE({sd} > DATE '{rd.isoformat()}', FALSE) AS POSTDATED
                FROM read_parquet('{cache_path.as_posix()}')
                WHERE DEALTYPE IN ({dt})
            )
            SELECT
                SUM(CASE WHEN COALESCE(CUSTFISS,'') NOT IN ({dfi}) AND NOT POSTDATED THEN AMTPAY END),
                SUM(CASE WHEN COALESCE(CUSTFISS,'') IN     ({dfi})                   THEN AMTPAY END),
                SUM(CASE WHEN COALESCE(CUSTFISS,'') NOT IN ({dfi}) AND POSTDATED     THEN AMTPAY END)
            FROM base
        """).fetchone()
    finally:
        con.close()
    return {k_main: r[0], k_dfi: r[1], k_pd: r[2]}


def _pmmd_raw(utfx_cache: Path, iutfx_cache: Path, rd: date) -> dict:
    out = {}
    out.update(_sum_deals(utfx_cache,  ('BCD', 'BCQ'),
                          rd, "TOTCPMMD", "TOTCDFI", "PDCPMMD"))
    out.update(_sum_deals(iutfx_cache, ('BCS', 'BCT', 'BCW', 'BQD'),
                          rd, "TOTIPMMD", "TOTIDFI", "PDIPMMD"))
    return out


def _pmmd_thousands(raw: dict) -> dict:
    """RM'000 figures for the 6 columns changed by eSMR 4294."""
    totcpmmd = _thousands(raw.get("TOTCPMMD"))
    totipmmd = _thousands(raw.get("TOTIPMMD"))
    pdcpmmd  = _thousands(raw.get("PDCPMMD"))
    pdipmmd  = _thousands(raw.get("PDIPMMD"))
    return {
        "TOTCPMMD": totcpmmd,
        "TOTIPMMD": totipmmd,
        "TOTNPMMD": _sas_sum(totcpmmd, totipmmd),
        "PDCPMMD":  pdcpmmd,
        "PDIPMMD":  pdipmmd,
        "PDNPMMD":  _sas_sum(pdcpmmd, pdipmmd),
    }


UTFX_CACHE  = _load_cached(INPUT_UTFX_FILE,  "UTFX")
IUTFX_CACHE = _load_cached(INPUT_IUTFX_FILE, "IUTFX")

_raw_today = _pmmd_raw(UTFX_CACHE, IUTFX_CACHE, reptdate)
_discard(UTFX_CACHE, IUTFX_CACHE)                  # temporary parquet no longer needed

utfx = pl.DataFrame(
    [{"REPTDATE": reptdate,
      "TOTCPMMD": _raw_today["TOTCPMMD"],
      "TOTCDFI":  _raw_today["TOTCDFI"],
      "PDCPMMD":  _raw_today["PDCPMMD"]}],
    schema={"REPTDATE": pl.Date, "TOTCPMMD": pl.Float64,
            "TOTCDFI": pl.Float64, "PDCPMMD": pl.Float64},
)
iutfx = pl.DataFrame(
    [{"REPTDATE": reptdate,
      "TOTIPMMD": _raw_today["TOTIPMMD"],
      "TOTIDFI":  _raw_today["TOTIDFI"],
      "PDIPMMD":  _raw_today["PDIPMMD"]}],
    schema={"REPTDATE": pl.Date, "TOTIPMMD": pl.Float64,
            "TOTIDFI": pl.Float64, "PDIPMMD": pl.Float64},
)
print(f"  UTFX : TOTCPMMD={_raw_today['TOTCPMMD']}  PDCPMMD={_raw_today['PDCPMMD']}  TOTCDFI={_raw_today['TOTCDFI']}")
print(f"  IUTFX: TOTIPMMD={_raw_today['TOTIPMMD']}  PDIPMMD={_raw_today['PDIPMMD']}  TOTIDFI={_raw_today['TOTIDFI']}")

# ============================================================================
# STEP 7: ERATE  (EFORATE)
# ============================================================================
print("\nStep 7: Reading ERATE (EFORATE)...")

EFORATE_CACHE = _load_cached(INPUT_EFORATE_FILE, "EFORATE")

con = duckdb.connect(database=":memory:")
erate = con.execute(f"""
    SELECT
        DATE '1960-01-01' + CAST(REPTDATE AS INTEGER) AS REPTDATE,
        CAST(SELLRATE AS DOUBLE) AS SELLRATE,
        CAST(BUYRATE  AS DOUBLE) AS BUYRATE
    FROM read_parquet('{EFORATE_CACHE.as_posix()}')
""").pl()
con.close()
_discard(EFORATE_CACHE)                            # temporary parquet no longer needed
print(f"  ERATE rows (whole month): {len(erate):,}")
erate_today = erate.filter(pl.col("REPTDATE") == reptdate)

# ============================================================================
# STEP 8: RNID / IRNID  (unchanged logic)
# ============================================================================
print("\nStep 8: Building RNID / IRNID...")

RNID_CACHE  = _load_cached(INPUT_RNID_FILE,  "RNID")
IRNID_CACHE = _load_cached(INPUT_IRNID_FILE, "IRNID")

con = duckdb.connect(database=":memory:")
rnid = con.execute(f"""
    WITH tagged AS (
        SELECT
            CASE WHEN CUSTCODE IN (77,78) THEN ABS(CAST(CURBAL AS DOUBLE)) END AS TOTINDV,
            CASE WHEN CUSTCODE NOT IN (77,78) THEN ABS(CAST(CURBAL AS DOUBLE)) END AS TOTNIND
        FROM read_parquet('{RNID_CACHE.as_posix()}')
        WHERE PRODUCT = 320
          AND NIDSTAT = 'N'
          AND CDSTAT  = 'A'
    )
    SELECT SUM(TOTINDV) AS TOTINDV, SUM(TOTNIND) AS TOTNIND
    FROM tagged
""").pl()
con.close()
rnid = rnid.with_columns(pl.lit(reptdate).alias("REPTDATE"))

con = duckdb.connect(database=":memory:")
irnid = con.execute(f"""
    WITH tagged AS (
        SELECT
            CASE WHEN CUSTCODE IN (77,78) THEN ABS(CAST(CURBAL AS DOUBLE)) END AS TOIINDV,
            CASE WHEN CUSTCODE NOT IN (77,78) THEN ABS(CAST(CURBAL AS DOUBLE)) END AS TOININD
        FROM read_parquet('{IRNID_CACHE.as_posix()}')
        WHERE PRODUCT = 321
          AND NIDSTAT = 'N'
          AND CDSTAT  = 'A'
    )
    SELECT SUM(TOIINDV) AS TOIINDV, SUM(TOININD) AS TOININD
    FROM tagged
""").pl()
con.close()
_discard(RNID_CACHE, IRNID_CACHE)                  # temporary parquet no longer needed
irnid = irnid.with_columns(pl.lit(reptdate).alias("REPTDATE"))

print(f"  RNID : TOTINDV={rnid['TOTINDV'][0]}  TOTNIND={rnid['TOTNIND'][0]}")
print(f"  IRNID: TOIINDV={irnid['TOIINDV'][0]}  TOININD={irnid['TOININD'][0]}")

# ============================================================================
# STEP 9: DCMG (report-date row)
# ============================================================================
print("\nStep 9: Building today's DCMG row...")

def _normalize_sas_reptdate(df: pl.DataFrame) -> pl.DataFrame:
    """Raw SAS numeric dates are days since 1960-01-01 (not 1970-01-01)."""
    dtype = df.schema["REPTDATE"]
    if dtype in (pl.Float32, pl.Float64, pl.Int32, pl.Int64):
        return df.with_columns(
            (pl.lit(date(1960, 1, 1)) + pl.duration(days=pl.col("REPTDATE").cast(pl.Int64)))
            .alias("REPTDATE")
        )
    return df.with_columns(pl.col("REPTDATE").cast(pl.Date))


def _left_join(base_df: pl.DataFrame, other_df: pl.DataFrame) -> pl.DataFrame:
    if other_df.is_empty():
        return base_df
    return base_df.join(other_df, on="REPTDATE", how="left")

dcmg_today = pl.DataFrame({"REPTDATE": [reptdate]}, schema={"REPTDATE": pl.Date})
for piece in (equtms, pbgia_today, erate_today, rnid, irnid, utfx, iutfx):
    dcmg_today = _left_join(dcmg_today, piece)

row = dcmg_today.row(0, named=True)

proceed   = _thousands(row.get("PROCEED"))
iproceed  = _thousands(row.get("IPROCEED"))
totprocd  = _sas_sum(proceed, iproceed)
iprocefi  = _thousands(row.get("IPROCEFI"))
closbal   = _thousands(row.get("CLOSBAL"))
totindv   = _thousands(row.get("TOTINDV"))
totnind   = _thousands(row.get("TOTNIND"))
totrnid   = _sas_sum(totindv, totnind)
toiindv   = _thousands(row.get("TOIINDV"))
toinind   = _thousands(row.get("TOININD"))
toirnid   = _sas_sum(toiindv, toinind)
totanid   = _sas_sum(totrnid, toirnid)
pmmd      = _pmmd_thousands(row)                 # TOTCPMMD/TOTIPMMD/TOTNPMMD + PDCPMMD/PDIPMMD/PDNPMMD
totcdfi   = _thousands(row.get("TOTCDFI"))
totidfi   = _thousands(row.get("TOTIDFI"))
totndfi   = _sas_sum(totcdfi, totidfi)

# Explicit schema so an all-missing column is never typed as Null
DCMG_SCHEMA = {
    "REPTDATE": pl.Date,
    "PROCEED": pl.Int64, "IPROCEED": pl.Int64, "TOTPROCD": pl.Int64, "IPROCEFI": pl.Int64,
    "TOTINDV": pl.Int64, "TOTNIND": pl.Int64, "TOTRNID": pl.Int64,
    "TOIINDV": pl.Int64, "TOININD": pl.Int64, "TOIRNID": pl.Int64, "TOTANID": pl.Int64,
    "TOTCPMMD": pl.Int64, "TOTIPMMD": pl.Int64, "TOTNPMMD": pl.Int64,
    "PDCPMMD": pl.Int64, "PDIPMMD": pl.Int64, "PDNPMMD": pl.Int64,
    "TOTCDFI": pl.Int64, "TOTIDFI": pl.Int64, "TOTNDFI": pl.Int64,
    "CLOSBAL": pl.Int64,
    "SELLRATE": pl.Float64, "BUYRATE": pl.Float64,
}

dcmg_today = pl.DataFrame([{
    "REPTDATE": reptdate,
    "PROCEED":  proceed,
    "IPROCEED": iproceed,
    "TOTPROCD": totprocd,
    "IPROCEFI": iprocefi,
    "TOTINDV":  totindv,
    "TOTNIND":  totnind,
    "TOTRNID":  totrnid,
    "TOIINDV":  toiindv,
    "TOININD":  toinind,
    "TOIRNID":  toirnid,
    "TOTANID":  totanid,
    "TOTCPMMD": pmmd["TOTCPMMD"],
    "TOTIPMMD": pmmd["TOTIPMMD"],
    "TOTNPMMD": pmmd["TOTNPMMD"],
    "PDCPMMD":  pmmd["PDCPMMD"],
    "PDIPMMD":  pmmd["PDIPMMD"],
    "PDNPMMD":  pmmd["PDNPMMD"],
    "TOTCDFI":  totcdfi,
    "TOTIDFI":  totidfi,
    "TOTNDFI":  totndfi,
    "CLOSBAL":  closbal,
    "SELLRATE": row.get("SELLRATE"),
    "BUYRATE":  row.get("BUYRATE"),
}], schema=DCMG_SCHEMA)
print("  Today's DCMG row built.")

# ============================================================================
# STEP 10: MONTHLY MASTER
#   Earlier rows come from the saved master (parquet), or from the legacy
#   .sas7bdat on the first run. REPTDAY = '01' restarts the month
#   (as in the SAS %APPEND macro).
# ============================================================================
print("\nStep 10: Updating monthly DCMG master...")

def _load_existing() -> pl.DataFrame:
    if STORE_DCMG_FILE.exists():
        return pl.read_parquet(STORE_DCMG_FILE)
    if INPUT_DCMG_LEGACY_FILE.exists():
        return _normalize_sas_reptdate(
            pl.read_parquet(_load_cached(INPUT_DCMG_LEGACY_FILE, "DCMG_LEGACY"))
        )
    return pl.DataFrame(schema=DCMG_SCHEMA)


def _stack_month() -> pl.DataFrame:
    """Today's row first, then earlier rows; NODUPKEY BY REPTDATE keeps today's."""
    if REPTDAY == "01":
        return dcmg_today
    existing = _load_existing()
    combined = pl.concat([dcmg_today, existing], how="diagonal_relaxed")
    return (
        combined
        .unique(subset=["REPTDATE"], keep="first", maintain_order=True)
        .sort("REPTDATE")
    )


# Unused-4 : old version without temporary-parquet clean-up (replaced by the version below)
# def _recompute_pmmd_for(d: date) -> Optional[dict]:
#     """Recompute the 6 PMMD columns of date d from that day's own UTFX / IUTFX."""
#     u_sas, i_sas = _utfx_file(d), _iutfx_file(d)
#     if not (u_sas.exists() and i_sas.exists()):
#         return None
#     u_cache = _load_cached(u_sas, f"UTFX {d:%y%m%d}")
#     i_cache = _load_cached(i_sas, f"IUTFX {d:%y%m%d}")
#     return _pmmd_thousands(_pmmd_raw(u_cache, i_cache, d))


def _recompute_pmmd_for(d: date) -> Optional[dict]:
    """Recompute the 6 PMMD columns of date d from that day's own UTFX / IUTFX."""
    u_sas, i_sas = _utfx_file(d), _iutfx_file(d)
    if not (u_sas.exists() and i_sas.exists()):
        return None
    u_cache = _load_cached(u_sas, f"UTFX {d:%y%m%d}")
    i_cache = _load_cached(i_sas, f"IUTFX {d:%y%m%d}")
    try:
        return _pmmd_thousands(_pmmd_raw(u_cache, i_cache, d))
    finally:
        _discard(u_cache, i_cache)           # one day at a time, never accumulates


def _recompute_earlier_rows(master: pl.DataFrame) -> pl.DataFrame:
    """Bring every earlier row of the month onto the same basis as the report-date row."""
    rows, missing = [], []
    for r in master.iter_rows(named=True):
        if r["REPTDATE"] < reptdate:
            fresh = _recompute_pmmd_for(r["REPTDATE"])
            if fresh is None:
                missing.append(r["REPTDATE"])
            else:
                r.update(fresh)
        rows.append(r)
    if missing:
        print("  WARNING: UTFX/IUTFX file missing, row left as saved for: "
              + ", ".join(d.strftime("%d/%m/%Y") for d in missing))
    return pl.DataFrame(rows, schema=master.schema, strict=False)


# ============================================================================
# STEP 11: REPORT DEFINITION  (one table; Excel + TXT)
# ============================================================================
print("\nStep 11: Report definition...")

# format_width = SAS format display width; on-page slot = format_width + 2 (TXT only)
REPORT_COLUMNS = [
    {"var": "REPTDATE", "header": ["DATE"],                                           "format_width": 10, "type": "date"},
    {"var": "PROCEED",  "header": ["NID", "(NON-FI)", "(PROCEEDS)", "(CONV)"],        "format_width": 15, "type": "comma"},
    {"var": "IPROCEED", "header": ["NID", "(NON-FI)", "(PROCEEDS)", "(ISL)"],         "format_width": 15, "type": "comma"},
    {"var": "TOTPROCD", "header": ["TOTAL", "NIDS", "(CONV + ISL)"],                  "format_width": 15, "type": "comma"},
    {"var": "IPROCEFI", "header": ["NID-I(FI)", "(PBB ONLY)", "(PROCEEDS)"],          "format_width": 15, "type": "comma"},
    {"var": "TOTINDV",  "header": ["INDV", "RETAIL", "NID", "(CONV)"],                "format_width": 15, "type": "comma"},
    {"var": "TOTNIND",  "header": ["NON-INDV", "RETAIL", "NID", "(CONV)"],            "format_width": 15, "type": "comma"},
    {"var": "TOTRNID",  "header": ["TOTAL", "RETAIL", "NID", "(CONV)"],               "format_width": 15, "type": "comma"},
    {"var": "TOIINDV",  "header": ["INDV", "RETAIL", "NIDC", "(ISL)"],                "format_width": 15, "type": "comma"},
    {"var": "TOININD",  "header": ["NON-INDV", "RETAIL", "NIDC", "(ISL)"],            "format_width": 15, "type": "comma"},
    {"var": "TOIRNID",  "header": ["TOTAL", "RETAIL", "NIDC", "(ISL)"],               "format_width": 15, "type": "comma"},
    {"var": "TOTANID",  "header": ["TOTAL", "RETAIL", "NID & NIDC", "(CONV + ISL)"],  "format_width": 15, "type": "comma"},
    {"var": "TOTCPMMD", "header": ["PMMD &", "PQMMD", "(CONV)"],                      "format_width": 15, "type": "comma"},
    {"var": "TOTIPMMD", "header": ["PMMD", "(ISL)"],                                  "format_width": 15, "type": "comma"},
    {"var": "TOTNPMMD", "header": ["TOTAL", "PMMD", "(CONV + ISL)"],                  "format_width": 15, "type": "comma"},
    {"var": "TOTCDFI",  "header": ["DFI & FI", "PMMD &", "PQMMD", "(CONV)"],          "format_width": 15, "type": "comma"},
    {"var": "TOTIDFI",  "header": ["DFI & FI", "PMMD", "(ISL)"],                      "format_width": 15, "type": "comma"},
    {"var": "TOTNDFI",  "header": ["TOTAL", "DFI & FI", "PMMD", "(CONV + ISL)"],      "format_width": 15, "type": "comma"},
    # --- NEW (eSMR 4294): post-dated PMMD, placed before GIA BALANCE ---
    {"var": "PDCPMMD",  "header": ["POST-DATED", "PMMD &", "PQMMD", "(CONV)"],        "format_width": 15, "type": "comma"},
    {"var": "PDIPMMD",  "header": ["POST-DATED", "PMMD", "(ISL)"],                    "format_width": 15, "type": "comma"},
    {"var": "PDNPMMD",  "header": ["TOTAL", "POST-DATED", "PMMD", "(CONV + ISL)"],    "format_width": 15, "type": "comma"},
    {"var": "CLOSBAL",  "header": ["GIA", "BALANCE", "IN RM", "(SPOTRATE)"],          "format_width": 15, "type": "comma"},
    {"var": "SELLRATE", "header": ["GIA", "CLSG PRICE", "RM/G", "(SELLING)"],         "format_width": 14, "type": "num2"},
    {"var": "BUYRATE",  "header": ["GIA", "CLSG PRICE", "RM/G", "(BUYING)"],          "format_width": 14, "type": "num2"},
]

HEADER_DEPTH = max(len(c["header"]) for c in REPORT_COLUMNS)


def _slot_width(col: dict) -> int:
    return col["format_width"] + 2


def _title_block() -> list[str]:
    return [
        "REPORT ID : DMMISR1G",
        "PUBLIC BANK BERHAD",
        "SALES ADMINISTRATION & SUPPORT",
        "BANK'S TOTAL DEPOSITS (RM '000)",
        f"AS AT {RPDATE}",
    ]


def _txt_header_block() -> list[str]:
    """Multi-line headers, right-justified, bottom-aligned, then the rule line."""
    lines = []
    for depth in range(HEADER_DEPTH):
        parts = []
        for col in REPORT_COLUMNS:
            hp  = col["header"]
            idx = depth - (HEADER_DEPTH - len(hp))
            parts.append((hp[idx] if idx >= 0 else "").rjust(_slot_width(col)))
        lines.append("".join(parts))
    total = sum(_slot_width(c) for c in REPORT_COLUMNS)
    lines.append("  " + "-" * (total - 2))
    return lines


def _txt_value(col: dict, value) -> str:
    width = _slot_width(col)
    if col["type"] == "date":
        return value.strftime("%d/%m/%Y").rjust(width) if value else "0".rjust(width)
    if col["type"] == "comma":
        return _fmt_comma(value, width, 0)
    return _fmt_num(value, width, 2)


def write_txt(df: pl.DataFrame, path: Path) -> None:
    """Single wide table: no column panels, no pagination."""
    lines = list(_title_block())
    lines.append("")
    lines.extend(_txt_header_block())
    for r in df.sort("REPTDATE").iter_rows(named=True):
        lines.append("".join(_txt_value(c, r.get(c["var"])) for c in REPORT_COLUMNS))
    with open(path, "w", encoding="latin1") as fh:
        fh.write("\n".join(lines) + "\n")


def write_excel(df: pl.DataFrame, path: Path) -> None:
    """Title block, then one header row and one table; real numeric cells, missing -> 0."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DMMISR1G"

    # --- title block (rows 1-5), one blank row, then the column header ---
    titles = _title_block()
    for i, t in enumerate(titles, start=1):
        ws.cell(row=i, column=1, value=t).font = Font(bold=True)
    hdr_row   = len(titles) + 2          # row 7 (row 6 stays blank)
    first_row = hdr_row + 1              # first data row = 8

    for j, col in enumerate(REPORT_COLUMNS, start=1):
        c = ws.cell(row=hdr_row, column=j, value="\n".join(col["header"]))
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.row_dimensions[hdr_row].height = 62

    for i, r in enumerate(df.sort("REPTDATE").iter_rows(named=True), start=first_row):
        for j, col in enumerate(REPORT_COLUMNS, start=1):
            v = r.get(col["var"])
            if col["type"] == "date":
                cell = ws.cell(row=i, column=j, value=v)
                cell.number_format = "DD/MM/YYYY"
            elif col["type"] == "comma":
                cell = ws.cell(row=i, column=j, value=(v if v is not None else 0))
                cell.number_format = "#,##0"
            else:
                cell = ws.cell(row=i, column=j, value=(v if v is not None else 0))
                cell.number_format = "#,##0.00"

    ws.freeze_panes = f"B{first_row}"
    wb.save(path)

# ============================================================================
# BUILD MONTH, SAVE MASTER, WRITE OUTPUTS
# ============================================================================
dcmg_monthly = _stack_month()
dcmg_monthly = _recompute_earlier_rows(dcmg_monthly)

dcmg_monthly.write_parquet(STORE_DCMG_FILE)
write_txt(dcmg_monthly, OUTPUT_TXT_FILE)
write_excel(dcmg_monthly, OUTPUT_XLSX_FILE)

print(f"\n  Master rows : {len(dcmg_monthly):,}")
print(f"  Master      : {STORE_DCMG_FILE}")
print(f"  TXT         : {OUTPUT_TXT_FILE}")
print(f"  Excel       : {OUTPUT_XLSX_FILE}")

del dcmg_today, dcmg_monthly
gc.collect()

# ============================================================================
# SFTP TO DRR  (JCL STEP01 + RUNSFTP)
# ============================================================================
SFTP_ENABLED      = True      # set False for UAT / reptdate override runs
SFTP_CONTROL_FILE = Path("/sasdata/dwh/control/ctl_dwh_sftp_info.sas7bdat")
DRR_HOST_DESC     = "Data Report Repository (DRR)"      # JCL: DRR#SFTP
DRR_REMOTE_DIR    = "WMTB-SA&S/DEP - SAS/BANK TOTAL"    # JCL STEP01: cd "..."


def get_sftp_info(host_desc: str):
    """Read one SFTP configuration row from the DWH control file (same as EIWFRMCR.py)."""
    if not SFTP_CONTROL_FILE.exists():
        raise FileNotFoundError(f"SFTP control file not found: {SFTP_CONTROL_FILE}")

    ctl = pd.read_sas(SFTP_CONTROL_FILE, format="sas7bdat", encoding="utf-8")
    ctl.columns = [str(c).upper().strip() for c in ctl.columns]

    required = {"HOST_DESC", "SFTP_ID", "SFTP_PW", "HOST_IP"}
    missing_columns = sorted(required.difference(ctl.columns))
    if missing_columns:
        raise RuntimeError("SFTP control file is missing columns: " + ", ".join(missing_columns))

    descriptions = ctl["HOST_DESC"].fillna("").astype(str).str.strip()
    rows = ctl[descriptions.eq(host_desc)]
    if len(rows) != 1:
        raise RuntimeError(f"Expected exactly one SFTP row for '{host_desc}'; found {len(rows)}")

    row = rows.iloc[0]
    clean = lambda value: "".join(str(value).split()) if pd.notna(value) else ""
    sftp_id  = clean(row["SFTP_ID"])
    sftp_pw  = clean(row["SFTP_PW"])
    host_ip  = clean(row["HOST_IP"])
    host_key = clean(row["HOST_KEY"]) if "HOST_KEY" in ctl.columns else ""

    empty_values = [n for n, v in (("SFTP_ID", sftp_id), ("SFTP_PW", sftp_pw), ("HOST_IP", host_ip)) if not v]
    if empty_values:
        raise RuntimeError(f"SFTP row for '{host_desc}' has blank values: {', '.join(empty_values)}")

    return sftp_id, sftp_pw, host_ip, host_key


def _drr_remote_filename() -> str:
    # JCL: put //SAP.PBB.DMMISR1G.DAILY  DMMISR1G@%OMM.%OYY..TXT
    return f"DMMISR1G@{REPTMON}{REPTYEAR}.TXT"


def sftp_upload_to_drr(local_path: Path) -> None:
    """
    Equivalent of the JCL:
        STEP01  : cd "WMTB-SA&S/DEP - SAS/BANK TOTAL"
        RUNSFTP : put //SAP.PBB.DMMISR1G.DAILY  DMMISR1G@%OMM.%OYY..TXT
    """
    if not local_path.exists():
        raise FileNotFoundError(f"Output file to SFTP not found: {local_path}")

    sftp_id, sftp_pw, host_ip, host_key = get_sftp_info(DRR_HOST_DESC)

    ssh = paramiko.SSHClient()
    # HOST_KEY exists in the control file but its format is unconfirmed;
    # AutoAddPolicy is used in the interim (same as EIWFRMCR.py).
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())

    print(f"  Connecting to DRR host {host_ip} as {sftp_id} ...")
    ssh.connect(hostname=host_ip, username=sftp_id, password=sftp_pw)
    sftp = ssh.open_sftp()
    try:
        try:
            sftp.chdir(DRR_REMOTE_DIR)
        except IOError:
            sys.exit(f"Aborting: Remote DRR folder '{DRR_REMOTE_DIR}' not found")

        remote_name = _drr_remote_filename()
        print(f"  SFTP upload: {local_path.name} -> {DRR_REMOTE_DIR}/{remote_name}")
        sftp.put(str(local_path), remote_name)

        # confirm the file landed complete
        if sftp.stat(remote_name).st_size != local_path.stat().st_size:
            raise RuntimeError(f"Size mismatch after upload of {remote_name}")
    finally:
        sftp.close()
        ssh.close()
    print("  DRR upload complete.")


if SFTP_ENABLED:
    print("\nUploading report to DRR via SFTP...")
    sftp_upload_to_drr(OUTPUT_TXT_FILE)

# ============================================================================
# JCL/HOUSEKEEPING STEPS RETAINED AS COMMENTS (no Python equivalent needed)
# ============================================================================
# //DELETE   EXEC PGM=IEFBR14  -> delete of SAP.PBB.DMMISR1G.DAILY (mainframe housekeeping)
# //RUNSFTP  EXEC COZBATCH     -> SFTP of DMMISR1G@%OMM.%OYY..TXT to DRR (operational, out of scope)
# //RUNSFTP  EXEC COZBATCH     -> NOW implemented above in sftp_upload_to_drr()

print("\nEIBDDCMG complete.")
